"""Send every student's weekly report to the parent, in one pass, after the teacher approves (M5).

Flow (the teacher is in charge of every step):
  1. Import  an Excel/CSV sheet: student, parent email, parent WhatsApp number -> saved as the parent contacts.
  2. Prepare one draft per student: the report's fixed Arabic text plus the share link. Nothing leaves the server.
  3. Review  the teacher reads and edits each draft, then approves it (or skips it).
  4. Send    approved drafts go out by email in one click. WhatsApp drafts open one by one in the teacher's own
             WhatsApp (wa.me), the teacher presses Send there, then marks it done.

A draft is only ever sent when its status is "approved". Re-preparing never touches a draft that was already sent.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any, Callable

from rafeeq.teacher import parent_delivery as pd

MAX_BYTES = 1_000_000
MAX_ROWS = 500
STATUSES = ("pending", "approved", "skipped")

# header aliases (lower-cased, spaces/underscores removed). Arabic and English headers both work.
_ALIASES = {
    "student": ("student", "studentname", "name", "الطالب", "اسمالطالب", "اسم", "الاسم"),
    "student_email": ("studentemail", "إيميلالطالب", "ايميلالطالب", "بريدالطالب"),
    "parent_email": ("parentemail", "emailparent", "parentmail", "email", "إيميلوليالأمر", "ايميلوليالامر",
                     "بريدوليالأمر", "بريدوليالامر", "الايميل", "الإيميل", "البريد"),
    "parent_phone": ("parentphone", "phone", "whatsapp", "mobile", "parentwhatsapp", "واتساب", "واتس",
                     "رقموليالأمر", "رقموليالامر", "رقمواتساب", "الموبايل", "رقمالموبايل", "الهاتف"),
}


def _fold(s: str) -> str:
    s = re.sub(r"[\u064b-\u065f\u0640]", "", s or "")           # tashkeel, tatweel
    return s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")


def _key(h: Any) -> str:
    return re.sub(r"[\s_\-]+", "", _fold(str(h or "")).strip().lower())


def _column_map(header: list) -> dict[str, int]:
    keys = [_key(h) for h in header]
    out: dict[str, int] = {}
    for field, names in _ALIASES.items():
        for name in names:
            if _key(name) in keys:
                out[field] = keys.index(_key(name))
                break
    # headers written in full ("رقم واتساب ولي الأمر", "Parent's WhatsApp number"): classify by keyword
    for i, k in enumerate(keys):
        if i in out.values() or not k:
            continue
        if re.search(r"ايميل|email|mail|بريد", k):
            field = "student_email" if re.search(r"طالب|student", k) and not re.search(r"ولي|parent|اهل", k) else "parent_email"
        elif re.search(r"واتس|whats|phone|mobile|موبايل|هاتف|تليفون|رقم", k):
            field = "parent_phone"
        elif re.search(r"طالب|student|اسم|name", k):
            field = "student"
        else:
            continue
        out.setdefault(field, i)
    return out


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():  # Excel stores 1012345678 as a number
        return str(int(v))
    return str(v).strip()


def parse_sheet(filename: str, data: bytes) -> list[dict]:
    """Rows of {student, student_email, parent_email, parent_phone} from .xlsx or .csv. Raises ValueError."""
    if len(data) > MAX_BYTES:
        raise ValueError("too_big")
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        except ImportError:
            raise ValueError("no_openpyxl")
        except Exception:
            raise ValueError("bad_file")
    else:
        raise ValueError("bad_type")
    rows = [r for r in rows if any(_cell(c) for c in r)]
    if not rows:
        raise ValueError("empty")
    cols = _column_map(rows[0])
    if "parent_email" not in cols and "parent_phone" not in cols or ("student" not in cols and "student_email" not in cols):
        raise ValueError("columns")
    out = []
    for r in rows[1: MAX_ROWS + 1]:
        item = {f: (_cell(r[i]) if i < len(r) else "") for f, i in cols.items()}
        for f in _ALIASES:
            item.setdefault(f, "")
        out.append(item)
    return out


def _norm_name(s: str) -> str:
    s = re.sub(r"[ً-ٟـ]", "", s or "")           # tashkeel, tatweel
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ى", "ي").replace("ة", "ه")
    return re.sub(r"\s+", " ", s).strip().lower()


def match_rows(rows: list[dict], students: list[dict]) -> tuple[list[tuple[dict, dict]], list[dict]]:
    """(matched [(student, row)], unmatched [{row, reason}]). Match by the student's email, else by name."""
    by_email = {s["email"].strip().lower(): s for s in students if s.get("email")}
    by_name: dict[str, list[dict]] = {}
    for s in students:
        by_name.setdefault(_norm_name(s.get("name", "")), []).append(s)
    matched, unmatched, used = [], [], set()
    for row in rows:
        s = by_email.get(row.get("student_email", "").strip().lower()) if row.get("student_email") else None
        if s is None and row.get("student"):
            cands = by_name.get(_norm_name(row["student"]), [])
            if len(cands) > 1:
                unmatched.append({"row": row, "reason": "ambiguous"})
                continue
            s = cands[0] if cands else None
        if s is None:
            unmatched.append({"row": row, "reason": "not_found"})
        elif s["user_id"] in used:
            unmatched.append({"row": row, "reason": "duplicate"})
        else:
            used.add(s["user_id"])
            matched.append((s, row))
    return matched, unmatched


def import_contacts(store: Any, room_id: str, students: list[dict], rows: list[dict]) -> dict:
    """Save the parent contacts from the sheet. Returns counts plus the rows that could not be used."""
    matched, unmatched = match_rows(rows, students)
    saved, invalid = 0, []
    for s, row in matched:
        try:
            pd.save_contact(store, room_id, s["user_id"], row.get("parent_email", ""), row.get("parent_phone", ""))
            saved += 1
        except ValueError as e:
            invalid.append({"row": row, "reason": str(e)})
    return {"saved": saved, "unmatched": unmatched, "invalid": invalid, "total": len(rows)}


# ---------------------------------------------------------------- drafts

def _ns(room_id: str) -> tuple:
    return ("class", room_id, "parent_drafts")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def get_draft(store: Any, room_id: str, user_id: str) -> dict | None:
    item = store.get(_ns(room_id), user_id)
    return dict(item.value) if item else None


def _put(store: Any, room_id: str, user_id: str, draft: dict) -> dict:
    store.put(_ns(room_id), user_id, draft)
    return draft


def list_drafts(store: Any, room_id: str) -> dict[str, dict]:
    out, offset = {}, 0
    while True:
        batch = store.search(_ns(room_id), limit=500, offset=offset)
        for it in batch:
            out[it.key] = dict(it.value)
        if len(batch) < 500:
            return out
        offset += 500


def prepare(store: Any, room_id: str, students: list[dict], report_for: Callable[[str], dict],
            link_for: Callable[[str], str]) -> dict:
    """One pending draft per student who has a parent contact. Drafts already sent are left alone."""
    made, skipped_no_contact, kept_sent = 0, [], 0
    for s in students:
        uid = s["user_id"]
        contact = pd.get_contact(store, room_id, uid)
        if not (contact["email"] or contact["phone"]):
            skipped_no_contact.append(s.get("name", uid))
            continue
        old = get_draft(store, room_id, uid) or {}
        if old.get("email_sent") or old.get("wa_sent"):
            kept_sent += 1
            continue
        subject, text = pd.compose(report_for(uid), link_for(uid))
        _put(store, room_id, uid, {"user_id": uid, "name": s.get("name", ""), "subject": subject, "text": text,
                                   "status": "pending", "email": contact["email"], "phone": contact["phone"],
                                   "email_sent": "", "wa_sent": "", "error": "", "prepared_at": _now()})
        made += 1
    return {"prepared": made, "no_contact": skipped_no_contact, "kept_sent": kept_sent}


class DraftError(Exception):
    pass


def edit_draft(store: Any, room_id: str, user_id: str, subject: str, text: str) -> dict:
    """The teacher changed the wording: it goes back to 'pending' and must be approved again."""
    d = get_draft(store, room_id, user_id)
    if d is None:
        raise DraftError("not_found")
    text = (text or "").strip()
    if not text:
        raise DraftError("empty")
    d.update(subject=(subject or "").strip()[:200] or d["subject"], text=text[:4000], status="pending")
    return _put(store, room_id, user_id, d)


def set_status(store: Any, room_id: str, user_id: str, status: str) -> dict:
    if status not in STATUSES:
        raise DraftError("bad_status")
    d = get_draft(store, room_id, user_id)
    if d is None:
        raise DraftError("not_found")
    d["status"] = status
    return _put(store, room_id, user_id, d)


def approve_all(store: Any, room_id: str) -> int:
    n = 0
    for uid, d in list_drafts(store, room_id).items():
        if d["status"] == "pending":
            d["status"] = "approved"
            _put(store, room_id, uid, d)
            n += 1
    return n


def mark_whatsapp_sent(store: Any, room_id: str, user_id: str) -> dict:
    d = get_draft(store, room_id, user_id)
    if d is None or d["status"] != "approved":
        raise DraftError("not_approved")
    d["wa_sent"] = _now()
    return _put(store, room_id, user_id, d)


def send_approved_emails(store: Any, room_id: str,
                         sender: Callable[[str, str, str], None] = pd.send_email) -> dict:
    """Email every approved draft that has an address and was not emailed yet. Stops early if email is off."""
    sent, failed = 0, []
    if not pd.email_configured():
        raise pd.DeliveryError("البريد غير مفعّل على هذا السيرفر. أضف BREVO_API_KEY وBREVO_SENDER في ملف .env.")
    for uid, d in list_drafts(store, room_id).items():
        if d["status"] != "approved" or not d.get("email") or d.get("email_sent"):
            continue
        try:
            sender(d["email"], d["subject"], d["text"])
        except pd.DeliveryError as e:
            d["error"] = str(e)
            failed.append({"name": d["name"], "error": str(e)})
        else:
            d["email_sent"], d["error"] = _now(), ""
            sent += 1
        _put(store, room_id, uid, d)
    return {"sent": sent, "failed": failed}


def summary(drafts: dict[str, dict]) -> dict:
    ds = list(drafts.values())
    return {"total": len(ds),
            "pending": sum(d["status"] == "pending" for d in ds),
            "approved": sum(d["status"] == "approved" for d in ds),
            "skipped": sum(d["status"] == "skipped" for d in ds),
            "emails_waiting": sum(d["status"] == "approved" and bool(d.get("email")) and not d.get("email_sent") for d in ds),
            "whatsapp_waiting": sum(d["status"] == "approved" and bool(d.get("phone")) and not d.get("wa_sent") for d in ds)}


# ---------------------------------------------------------------- the sheet, pre-filled for the teacher

TEMPLATE_HEAD = ("اسم الطالب", "إيميل الطالب", "إيميل ولي الأمر", "رقم واتساب ولي الأمر")


def build_template(store: Any, room_id: str, students: list[dict]) -> bytes:
    """An .xlsx with one row per student of the room (name, student email) and any parent details already saved.

    The teacher fills in or corrects the parent columns and uploads the file again. Student emails make the
    match exact. Phone numbers are written as text so the leading 0 / + is kept.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "أولياء الأمور"
    ws.sheet_view.rightToLeft = True
    ws.append(list(TEMPLATE_HEAD))
    for s in sorted(students, key=lambda x: _norm_name(x.get("name", ""))):
        c = pd.get_contact(store, room_id, s["user_id"])
        ws.append([s.get("name", ""), s.get("email", ""), c["email"], ("+" + c["phone"]) if c["phone"] else ""])
    thin = Side(style="thin", color="D3DDD8")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if cell.row == 1:
                cell.font = Font(bold=True, color="FFFFFF", size=12, name="Arial")
                cell.fill = PatternFill("solid", fgColor="087F73")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = Font(name="Arial", size=11, color="8A9A9C" if cell.column <= 2 else "202D32")
                cell.alignment = Alignment(horizontal="right" if cell.column == 1 else "left", vertical="center")
            if cell.column >= 2:
                cell.number_format = "@"
    ws.row_dimensions[1].height = 26
    for col, w in zip("ABCD", (30, 30, 32, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    note = wb.create_sheet("تعليمات")
    note.sheet_view.rightToLeft = True
    for line in ("كيف تملأ الملف (الورقة الأولى فقط تُقرأ):",
                 "• الأسماء وإيميلات الطلاب (الرمادية) من غرفتك، لا تغيّرها.",
                 "• املأ إيميل ولي الأمر و/أو رقم الواتساب لكل طالب (يكفي واحد). مثال الرقم: 01012345678.",
                 "• يمكنك حذف سطر أي طالب لا تريد إرسال تقريره.",
                 "• احفظ الملف .xlsx وارفعه من صفحة «إرسال التقارير لكل الأهالي»."):
        note.append([line])
    note.column_dimensions["A"].width = 90
    note["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
