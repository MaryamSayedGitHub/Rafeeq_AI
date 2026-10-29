"""The teacher studios as separate, connected pages: Presentation Studio (decks, themes, AI pictures),
Visual & Media Studio, My Learning Platform and the library that ties them together.

Runs against the REAL app (real routes, real stores, real files in a temp folder). Providers that cost
money or need the internet (Wan picture model, edge-tts, YouTube thumbnails) are replaced by small fakes.
    uv run pytest tests/test_studios.py
"""
import io
import re
import zipfile
from types import SimpleNamespace

import pytest

MONA = {"Authorization": "Bearer demo-token-mona-456"}   # teacher u_mona
SARA = {"Authorization": "Bearer demo-token-sara-123"}   # student u_sara
HX = {"HX-Request": "true"}


def png_bytes(w=240, h=120) -> bytes:
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    img = np.full((h, w, 3), 200, np.uint8)
    cv2.rectangle(img, (10, 10), (60, 60), (0, 0, 255), -1)
    return cv2.imencode(".png", img)[1].tobytes()


class FakeStore:
    """Same calls as the LangGraph store the app uses (put / get / search / delete)."""

    def __init__(self):
        self.data = {}

    def put(self, namespace, key, value):
        self.data[(tuple(namespace), key)] = dict(value)

    def get(self, namespace, key):
        value = self.data.get((tuple(namespace), key))
        return SimpleNamespace(key=key, value=value) if value is not None else None

    def search(self, namespace, limit=10):
        return [SimpleNamespace(key=k, value=v) for (ns, k), v in self.data.items() if ns == tuple(namespace)][:limit]

    def delete(self, namespace, key):
        self.data.pop((tuple(namespace), key), None)


@pytest.fixture
def folders(tmp_path, monkeypatch):
    for name in ("ASSET_DIR", "OUTPUT_DIR"):
        folder = (tmp_path / name).resolve()
        folder.mkdir()
        monkeypatch.setattr(f"rafeeq.teacher.lecturer.settings.{name}", folder)
    return tmp_path


@pytest.fixture
def store(folders):
    return FakeStore()


@pytest.fixture
def client(monkeypatch, folders):
    pytest.importorskip("langgraph")
    from fastapi.testclient import TestClient

    from rafeeq.platform import service as platform

    monkeypatch.setattr("rafeeq.teacher.lecturer.settings.DATABASE_URL", None)   # in-memory store
    monkeypatch.setenv("RAFEEQ_WEB_SECRET", "test-secret")
    avatars = (folders / "AVATARS").resolve()
    avatars.mkdir()
    monkeypatch.setattr("rafeeq.teacher.routes.AVATAR_DIR", avatars)
    platform.set_store(platform.PlatformStore(folders / "platform", seed_demo=True))
    from rafeeq.main import app

    with TestClient(app) as c:
        yield c
    platform.set_store(None)


def teacher_store(client):
    return client.app.state.teacher.store


def new_account(client, role, email):
    res = client.post("/auth/register", json={"name": f"{role} {email}", "email": email, "password": "secret1", "role": role})
    assert res.status_code in (200, 201), res.text
    return {"Authorization": f"Bearer {res.json()['token']}"}


def upload(client, headers=MONA, name="p.png", data=None, ctype="image/png", **form):
    res = client.post("/teacher/media/upload", headers=headers, files={"file": (name, data or png_bytes(), ctype)}, data=form)
    assert res.status_code == 201, res.text
    return res.json()


def make_deck(client, owner="u_mona", session_id="s_1", room_id=""):
    from rafeeq.teacher import decks

    return decks.create_deck(
        teacher_store(client), owner, title="الخلية", course_name="Biology", lecturer="Mona", room_id=room_id, session_id=session_id,
        slides=[{"title": "مقدمة", "bullets": ["أ", "ب"], "notes": "ن"}, {"title": "الخاتمة", "bullets": ["ج"]}])


# =============================================================== pure modules
def test_themes_validate_colours_and_extract_from_a_real_pptx():
    from rafeeq.teacher import themes

    ok = themes.normalize({"name": " Mine ", "font": "Cairo", "colors": {"blue": "#abcdef"}})
    assert ok["colors"]["blue"] == "#ABCDEF" and ok["colors"]["ink"] == themes.PRESETS["rafeeq"]["colors"]["ink"]
    for bad in ({"colors": {"blue": "red"}}, {"colors": {"ink": "#12345"}}, "nope"):
        with pytest.raises(themes.ThemeError):
            themes.normalize(bad)

    pptx = pytest.importorskip("pptx")
    buf = io.BytesIO()
    pptx.Presentation().save(buf)
    got = themes.extract_from_pptx(buf.getvalue(), fallback_name="Default")
    assert set(got["colors"]) == set(themes.COLOR_KEYS) and all(re.fullmatch(r"#[0-9A-F]{6}", c) for c in got["colors"].values())
    with pytest.raises(themes.ThemeError):
        themes.extract_from_pptx(b"not a zip")
    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(themes.ThemeError):
        themes.extract_from_pptx(empty.getvalue())


def test_image_edits_change_the_picture_and_refuse_bad_values():
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    from rafeeq.teacher import media_tools

    def shape(data):
        return cv2.imdecode(np.frombuffer(data, dtype=np.uint8), 1).shape[:2]

    src = png_bytes(240, 120)
    assert shape(media_tools.edit_image(src, ".png", {"rotate": 90})) == (240, 120)
    assert shape(media_tools.edit_image(src, ".png", {"width": 120})) == (60, 120)
    assert shape(media_tools.edit_image(src, ".png", {"crop": {"x": 0, "y": 0, "w": 50, "h": 50}})) == (60, 120)
    gray = cv2.imdecode(np.frombuffer(media_tools.edit_image(src, ".png", {"grayscale": True}), dtype=np.uint8), 1)
    assert (gray[..., 0] == gray[..., 1]).all()
    for bad in ({"rotate": 45}, {"flip": "x"}, {"crop": {"x": 90, "y": 0, "w": 50, "h": 50}}, {"brightness": 500}):
        with pytest.raises(media_tools.MediaError):
            media_tools.edit_image(src, ".png", bad)
    with pytest.raises(media_tools.MediaError):
        media_tools.edit_image(b"not an image", ".png", {})
    assert [media_tools.kind_for(n) for n in ("a.JPG", "b.mp4", "c.mp3", "d.exe")] == ["image", "video", "audio", ""]


def test_decks_are_per_teacher_idempotent_and_only_accept_own_images(store):
    from rafeeq.teacher import assets, decks

    slides = [{"title": "One", "bullets": ["a"]}]
    d = decks.create_deck(store, "u_mona", title="T", slides=slides, session_id="s1")
    assert decks.create_deck(store, "u_mona", title="T", slides=slides, session_id="s1")["deck_id"] == d["deck_id"]
    assert decks.get_deck(store, "u_other", d["deck_id"]) is None and decks.list_decks(store, "u_other") == []
    with pytest.raises(decks.DeckError):
        decks.create_deck(store, "u_mona", title="T", slides=[{"title": "", "bullets": []}])
    with pytest.raises(decks.DeckError):
        decks.update_deck(store, "u_mona", d["deck_id"], theme_id="no-such-theme")
    other_image = assets.register_asset(store, "u_other", kind="image", title="x", origin="upload")
    with pytest.raises(decks.DeckError):
        decks.set_slide_image(store, "u_mona", d["deck_id"], 0, other_image["asset_id"])
    with pytest.raises(decks.DeckError):
        decks.set_slide_image(store, "u_mona", d["deck_id"], 5, "")
    mine = assets.register_asset(store, "u_mona", kind="image", title="mine", origin="upload")
    assert decks.set_slide_image(store, "u_mona", d["deck_id"], 0, mine["asset_id"])["slides"][0]["image_asset_id"] == mine["asset_id"]
    assert decks.delete_deck(store, "u_mona", d["deck_id"]) and decks.get_deck(store, "u_mona", d["deck_id"]) is None


def test_courses_snapshot_only_what_is_ready_and_checks_ownership(store):
    from rafeeq.teacher import assets, courses

    ready = assets.register_asset(store, "u_mona", kind="source_video", title="ok", origin="external", source_url="https://youtu.be/x")
    pending = assets.register_asset(store, "u_mona", kind="image", title="later", origin="upload", status="pending")
    foreign = assets.register_asset(store, "u_other", kind="image", title="theirs", origin="upload")
    c = courses.create_course(store, "u_mona", "Biology")
    with pytest.raises(courses.CourseError):
        courses.create_course(store, "u_mona", "x")
    with pytest.raises(courses.CourseError):
        courses.add_lesson(store, "u_mona", c["course_id"], "L", asset_ids=[foreign["asset_id"]])
    with pytest.raises(courses.CourseError):
        courses.publish(store, "u_mona", c["course_id"], "r1")            # no lessons yet
    c = courses.add_lesson(store, "u_mona", c["course_id"], "Lesson 1", asset_ids=[ready["asset_id"], pending["asset_id"]])
    lesson = c["lessons"][0]["lesson_id"]
    published = courses.publish(store, "u_mona", c["course_id"], "r1")
    assert published["published_room_id"] == "r1" and not courses.has_unpublished_changes(published)
    snap = courses.published_for_room(store, "r1")
    assert [a["asset_id"] for a in snap[0]["lessons"][0]["assets"]] == [ready["asset_id"]]   # the pending one is left out
    assert courses.published_owner(store, "r1", c["course_id"], ready["asset_id"]) == "u_mona"
    assert courses.published_owner(store, "r1", c["course_id"], pending["asset_id"]) == ""
    assert courses.published_owner(store, "r2", c["course_id"], ready["asset_id"]) == ""
    edited = courses.update_lesson(store, "u_mona", c["course_id"], lesson, title="Renamed")
    assert courses.has_unpublished_changes(edited)
    courses.unpublish(store, "u_mona", c["course_id"])
    assert courses.published_for_room(store, "r1") == []
    assert courses.delete_course(store, "u_mona", c["course_id"]) and courses.get_course(store, "u_mona", c["course_id"]) is None


# =============================================================== API: themes, decks, export
def test_the_studio_api_needs_a_teacher_token(client):
    for path in ("/teacher/decks", "/teacher/themes", "/teacher/courses", "/teacher/media/capabilities"):
        assert client.get(path).status_code == 401, path
        assert client.get(path, headers=SARA).status_code == 403, path


def test_extract_a_theme_from_a_pptx_and_list_it(client):
    pptx = pytest.importorskip("pptx")
    buf = io.BytesIO()
    pptx.Presentation().save(buf)
    files = {"file": ("brand.pptx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res = client.post("/teacher/themes/extract", headers=MONA, files=files, data={"name": "My brand"})
    assert res.status_code == 201 and res.json()["kind"] == "theme" and res.json()["theme"]["name"] == "My brand"
    again = client.post("/teacher/themes/extract", headers=MONA, files=files, data={"name": "My brand"})
    assert again.json()["asset_id"] == res.json()["asset_id"]                              # same file, no duplicate
    listing = client.get("/teacher/themes", headers=MONA).json()
    assert [p["theme_id"] for p in listing["presets"]][0].startswith("preset:")
    assert [c["theme_id"] for c in listing["custom"]] == [res.json()["asset_id"]]
    assert client.get("/teacher/themes", headers=new_account(client, "teacher", "t2@x.com")).json()["custom"] == []
    bad = client.post("/teacher/themes/extract", headers=MONA, files={"file": ("x.pptx", b"nope", "application/zip")})
    assert bad.status_code == 422


def test_edit_a_deck_and_export_it_with_a_theme_and_a_picture(client):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = make_deck(client)
    other = new_account(client, "teacher", "t2@x.com")
    assert client.get(f"/teacher/decks/{deck['deck_id']}", headers=other).status_code == 404       # not theirs
    assert client.get("/teacher/decks", headers=other).json()["decks"] == []

    listing = client.get("/teacher/decks", headers=MONA).json()["decks"]
    assert [d["deck_id"] for d in listing] == [deck["deck_id"]] and listing[0]["slide_count"] == 2

    from rafeeq.teacher import assets, themes
    custom = themes.normalize({"name": "Green", "colors": {"paper": "#EAF4EE", "blue": "#087F73"}})
    theme_asset = assets.register_asset(teacher_store(client), "u_mona", kind="theme", title="Green", origin="upload", settings_used=custom)
    picture = upload(client)

    slides = [{"title": "مقدمة معدّلة", "bullets": ["نقطة 1", "نقطة 2"], "notes": "ملاحظة", "image_asset_id": picture["asset_id"]},
              {"title": "الخاتمة", "bullets": ["ج"]}]
    saved = client.put(f"/teacher/decks/{deck['deck_id']}", headers=MONA,
                       json={"title": "الخلية 2", "slides": slides, "theme_id": theme_asset["asset_id"]})
    assert saved.status_code == 200, saved.text
    view = saved.json()
    assert view["title"] == "الخلية 2" and view["theme_name"] == "Green" and view["theme"]["colors"]["paper"] == "#EAF4EE"
    assert view["slides"][0]["image_file_url"] == f"/teacher/assets/{picture['asset_id']}/file" and view["slides"][1]["image_file_url"] is None

    # refusals
    assert client.put(f"/teacher/decks/{deck['deck_id']}", headers=MONA, json={"theme_id": "preset:nope"}).status_code == 422
    assert client.put(f"/teacher/decks/{deck['deck_id']}", headers=MONA, json={"slides": []}).status_code == 422
    foreign = upload(client, headers=other)
    assert client.post(f"/teacher/decks/{deck['deck_id']}/slides/1/image", headers=MONA,
                       json={"asset_id": foreign["asset_id"]}).status_code == 422                      # someone else's image
    assert client.post(f"/teacher/decks/{deck['deck_id']}/slides/9/image", headers=MONA, json={"asset_id": ""}).status_code == 422

    exported = client.post(f"/teacher/decks/{deck['deck_id']}/export", headers=MONA)
    assert exported.status_code == 200, exported.text
    asset = exported.json()
    assert asset["kind"] == "pptx" and asset["origin"] == "edited" and asset["status"] == "ready" and asset["file_url"]
    data = client.get(asset["file_url"], headers=MONA).content
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 2
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sl in prs.slides for sh in sl.shapes)       # the picture is in the file
    fills = set()
    for sl in prs.slides:
        for sh in sl.shapes:
            try:
                if sh.fill.type == 1:
                    fills.add(str(sh.fill.fore_color.rgb))
            except (AttributeError, TypeError, ValueError):
                pass
    assert "EAF4EE" in fills                                                                          # the theme's paper colour

    again = client.post(f"/teacher/decks/{deck['deck_id']}/export", headers=MONA).json()
    assert again["version"] == asset["version"] + 1 and again["parent_asset_id"] == asset["asset_id"]  # versions, not overwrites
    assert client.get(asset["file_url"], headers=MONA).content == data                                 # the first file is untouched
    assert client.get(f"/teacher/decks/{deck['deck_id']}", headers=MONA).json()["exports"] == 2


# =============================================================== API: AI pictures for slides
def fake_picture_provider(monkeypatch, status="SUCCEEDED"):
    calls = SimpleNamespace(submitted=[], polled=[])
    monkeypatch.setattr("rafeeq.teacher.visuals.suggest_prompt", lambda title, bullets: (f"diagram of {title}", True))

    def submit(prompt):
        calls.submitted.append(prompt)
        return {"ok": True, "task_id": f"T{len(calls.submitted)}", "model": "wan-test"}

    def poll(task_id):
        calls.polled.append(task_id)
        return {"status": status, "image_url": "https://example.com/x.png" if status == "SUCCEEDED" else None,
                "error": None if status == "SUCCEEDED" else "blocked"}

    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.submit_task", submit)
    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.poll_task", poll)
    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.download_image", lambda url: png_bytes())
    return calls


def test_a_generated_slide_picture_lands_in_the_library_and_on_the_slide(client, monkeypatch):
    calls = fake_picture_provider(monkeypatch)
    deck = make_deck(client)
    url = f"/teacher/decks/{deck['deck_id']}/slides/0/visual"
    started = client.post(url, headers=MONA, json={})
    assert started.status_code == 200, started.text
    job = started.json()
    assert job["task_id"] == "T1" and job["prompt"] == "diagram of مقدمة" and job["used_llm"] is True
    assert client.post(url, headers=MONA, json={}).json()["asset_id"] == job["asset_id"] and len(calls.submitted) == 1   # a retry costs nothing

    status = client.get("/teacher/visuals/T1", headers=MONA).json()
    assert status["status"] == "SUCCEEDED" and status["file_url"] and status["asset_id"] == job["asset_id"]
    assert client.get("/teacher/visuals/T1", headers=MONA).json()["status"] == "SUCCEEDED" and calls.polled == ["T1"]    # not polled again
    asset = client.get(f"/teacher/assets/{job['asset_id']}", headers=MONA).json()
    assert asset["kind"] == "image" and asset["origin"] == "ai_generated" and asset["status"] == "ready" and asset["slide_index"] == 0
    assert client.get(f"/teacher/decks/{deck['deck_id']}", headers=MONA).json()["slides"][0]["image_asset_id"] == job["asset_id"]
    assert client.get("/teacher/visuals/T1", headers=new_account(client, "teacher", "t2@x.com")).status_code == 404


def test_a_failed_picture_job_is_reported_and_recorded(client, monkeypatch):
    fake_picture_provider(monkeypatch, status="FAILED")
    deck = make_deck(client)
    job = client.post(f"/teacher/decks/{deck['deck_id']}/slides/1/visual", headers=MONA, json={"prompt": "a red cell"}).json()
    assert job["prompt"] == "a red cell" and job["used_llm"] is False
    status = client.get(f"/teacher/visuals/{job['task_id']}", headers=MONA).json()
    assert status["status"] == "FAILED" and status["error"] == "blocked"
    assert client.get(f"/teacher/assets/{job['asset_id']}", headers=MONA).json()["status"] == "failed"
    assert client.post(f"/teacher/decks/{deck['deck_id']}/slides/7/visual", headers=MONA, json={}).status_code == 404


def test_picture_generation_without_a_provider_key_is_a_clear_503(client, monkeypatch):
    monkeypatch.setattr("rafeeq.teacher.visuals.suggest_prompt", lambda t, b: ("p", False))
    monkeypatch.setattr("rafeeq.teacher.lecturer.image_gen.submit_task",
                        lambda prompt: {"ok": False, "error": "DASHSCOPE_API_KEY is not set in .env"})
    deck = make_deck(client)
    res = client.post(f"/teacher/decks/{deck['deck_id']}/slides/0/visual", headers=MONA, json={})
    assert res.status_code == 503 and "DASHSCOPE_API_KEY" in res.json()["detail"]


# =============================================================== API: media
def test_upload_and_edit_keep_the_original_and_create_versions(client):
    cv2 = pytest.importorskip("cv2")
    np = pytest.importorskip("numpy")
    original = png_bytes()
    first = upload(client, data=original, title="Cell photo")
    assert first["kind"] == "image" and first["origin"] == "upload" and first["status"] == "ready" and first["title"] == "Cell photo"
    assert upload(client, data=original)["asset_id"] == first["asset_id"]                       # same file twice = one record

    edited = client.post(f"/teacher/media/{first['asset_id']}/edit", headers=MONA, json={"ops": {"rotate": 90, "grayscale": True}})
    assert edited.status_code == 201, edited.text
    new = edited.json()
    assert new["asset_id"] != first["asset_id"] and new["version"] == 2 and new["parent_asset_id"] == first["asset_id"]
    assert new["origin"] == "edited" and new["status"] == "ready"
    pic = cv2.imdecode(np.frombuffer(client.get(new["file_url"], headers=MONA).content, dtype=np.uint8), 1)
    assert pic.shape[:2] == (240, 120)
    assert client.get(first["file_url"], headers=MONA).content == original                       # original untouched
    again = client.post(f"/teacher/media/{first['asset_id']}/edit", headers=MONA, json={"ops": {"rotate": 90, "grayscale": True}})
    assert again.json()["asset_id"] == new["asset_id"]                                           # a double click makes no duplicate

    for ops in ({"rotate": 45}, {"crop": {"x": 90, "y": 0, "w": 50, "h": 50}}):
        assert client.post(f"/teacher/media/{first['asset_id']}/edit", headers=MONA, json={"ops": ops}).status_code == 422
    other = new_account(client, "teacher", "t2@x.com")
    assert client.post(f"/teacher/media/{first['asset_id']}/edit", headers=other, json={"ops": {"rotate": 90}}).status_code == 404


def test_upload_refuses_unknown_types_and_empty_files(client):
    res = client.post("/teacher/media/upload", headers=MONA, files={"file": ("run.exe", b"MZ", "application/octet-stream")})
    assert res.status_code == 422
    res = client.post("/teacher/media/upload", headers=MONA, files={"file": ("empty.png", b"", "image/png")})
    assert res.status_code in (413, 422)


def test_video_editing_says_so_when_ffmpeg_is_missing(client, monkeypatch):
    monkeypatch.setattr("rafeeq.teacher.media_tools.ffmpeg_path", lambda: None)
    clip = upload(client, name="clip.mp4", data=b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 64, ctype="video/mp4")
    assert clip["kind"] == "video"
    res = client.post(f"/teacher/media/{clip['asset_id']}/edit", headers=MONA, json={"ops": {"trim": {"start": 0, "end": 2}}})
    assert res.status_code == 503 and "ffmpeg" in res.json()["detail"]
    assert client.get("/teacher/media/capabilities", headers=MONA).json()["ffmpeg"] is False


def test_narration_becomes_an_audio_asset(client, monkeypatch):
    def fake_speech(text, voice, dst):
        dst.write_bytes(b"ID3-fake-mp3")
        return dst

    from rafeeq.teacher import media_tools

    real_speech = media_tools.synthesize_speech
    monkeypatch.setattr(media_tools, "synthesize_speech", fake_speech)
    body = {"text": "مرحبًا بكم في درس الخلية", "voice": "ar-EG-SalmaNeural", "title": "مقدمة صوتية"}
    res = client.post("/teacher/media/speech", headers=MONA, json=body)
    assert res.status_code == 201, res.text
    audio = res.json()
    assert audio["kind"] == "audio" and audio["origin"] == "ai_generated" and audio["status"] == "ready"
    assert client.get(audio["file_url"], headers=MONA).content == b"ID3-fake-mp3"
    assert client.post("/teacher/media/speech", headers=MONA, json=body).json()["asset_id"] == audio["asset_id"]
    monkeypatch.setattr(media_tools, "synthesize_speech", real_speech)
    assert client.post("/teacher/media/speech", headers=MONA, json={**body, "voice": "xx"}).status_code == 422


def test_youtube_cover_is_saved_from_youtube_only(client, monkeypatch):
    seen = []

    class Fake:
        def __init__(self, data):
            self.data = data

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self, n=-1):
            return self.data if n < 0 else self.data[:n]

    def fake_open(req, timeout=0):
        seen.append(req.full_url)
        return Fake(png_bytes())

    monkeypatch.setattr("urllib.request.urlopen", fake_open)
    res = client.post("/teacher/media/thumbnail", headers=MONA, json={"video_id": "dQw4w9WgXcQ", "title": "Cells", "start_sec": 75})
    assert res.status_code == 201, res.text
    cover = res.json()
    assert cover["kind"] == "image" and cover["origin"] == "external" and cover["status"] == "ready"
    assert cover["source_url"] == "https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=75s"
    assert seen == ["https://i.ytimg.com/vi/dQw4w9WgXcQ/hqdefault.jpg"]
    assert client.post("/teacher/media/thumbnail", headers=MONA, json={"video_id": "dQw4w9WgXcQ"}).json()["asset_id"] == cover["asset_id"]
    assert len(seen) == 1
    for bad in ("../../etc/passwd", "short", "https://evil.example/x.jpg"):
        assert client.post("/teacher/media/thumbnail", headers=MONA, json={"video_id": bad}).status_code == 422
    assert len(seen) == 1


# =============================================================== API: courses and what students see
def room_with_student(client):
    made = client.post("/rooms", headers=MONA, json={"name": "Biology 10"}).json()
    room_id, code = made["room"]["room_id"], made["join_code"]
    assert client.post("/rooms/join", headers=SARA, json={"code": code}).status_code == 200
    return room_id


def test_courses_are_published_to_the_room_and_students_see_only_the_snapshot(client):
    room_id = room_with_student(client)
    image, later = upload(client), upload(client, data=png_bytes(300, 150))
    course = client.post("/teacher/courses", headers=MONA, json={"title": "الأحياء", "room_id": room_id}).json()
    cid = course["course_id"]
    assert client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id}).status_code == 422   # empty course
    with_lesson = client.post(f"/teacher/courses/{cid}/lessons", headers=MONA,
                              json={"title": "الدرس الأول", "notes": "ابدأ هنا", "asset_ids": [image["asset_id"]]})
    assert with_lesson.status_code == 201 and with_lesson.json()["lessons"][0]["assets"][0]["asset_id"] == image["asset_id"]
    lid = with_lesson.json()["lessons"][0]["lesson_id"]

    assert client.get(f"/rooms/{room_id}/courses", headers=SARA).json() == {"courses": []}                                # not published yet
    published = client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    assert published.status_code == 200 and published.json()["published_room_id"] == room_id
    assert published.json()["unpublished_changes"] is False

    seen = client.get(f"/rooms/{room_id}/courses", headers=SARA).json()["courses"]
    assert [c["title"] for c in seen] == ["الأحياء"] and seen[0]["lessons"][0]["notes"] == "ابدأ هنا"
    asset = seen[0]["lessons"][0]["assets"][0]
    assert asset["file_url"] == f"/rooms/{room_id}/courses/{cid}/assets/{image['asset_id']}/file"
    got = client.get(asset["file_url"], headers=SARA)
    assert got.status_code == 200 and got.content == client.get(image["file_url"], headers=MONA).content

    # edits after publishing are invisible to students until the teacher publishes again
    after = client.post(f"/teacher/courses/{cid}/lessons/{lid}/assets", headers=MONA, json={"add": later["asset_id"]}).json()
    assert after["unpublished_changes"] is True
    assert client.get(f"/rooms/{room_id}/courses/{cid}/assets/{later['asset_id']}/file", headers=SARA).status_code == 404
    client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    assert client.get(f"/rooms/{room_id}/courses/{cid}/assets/{later['asset_id']}/file", headers=SARA).status_code == 200

    client.post(f"/teacher/courses/{cid}/unpublish", headers=MONA)
    assert client.get(f"/rooms/{room_id}/courses", headers=SARA).json() == {"courses": []}
    assert client.get(asset["file_url"], headers=SARA).status_code == 404


def test_course_progress_is_per_student_and_only_counts_real_items(client):
    room_id = room_with_student(client)
    image = upload(client)
    cid = client.post("/teacher/courses", headers=MONA, json={"title": "Progress", "room_id": room_id}).json()["course_id"]
    course = client.post(f"/teacher/courses/{cid}/lessons", headers=MONA,
                         json={"title": "L1", "notes": "Read me", "asset_ids": [image["asset_id"]]}).json()
    lid = course["lessons"][0]["lesson_id"]
    assert client.get(f"/rooms/{room_id}/courses/{cid}/progress", headers=SARA).status_code == 404     # not published yet
    client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    first = client.get(f"/rooms/{room_id}/courses/{cid}/progress", headers=SARA).json()
    assert first["done"] == [] and first["total"] == 2 and first["percent"] == 0
    ignored = client.post(f"/rooms/{room_id}/courses/{cid}/progress", headers=SARA, json={"item_id": "l_x:a_x"}).json()
    assert ignored["done"] == []
    half = client.post(f"/rooms/{room_id}/courses/{cid}/progress", headers=SARA, json={"item_id": f"{lid}:notes"}).json()
    assert half["percent"] == 50 and half["last"] == f"{lid}:notes"
    full = client.post(f"/rooms/{room_id}/courses/{cid}/progress", headers=SARA, json={"item_id": f"{lid}:{image['asset_id']}"}).json()
    assert full["percent"] == 100
    summary = client.get(f"/teacher/courses/{cid}/progress", headers=MONA).json()
    assert summary["started"] == 1 and summary["completed"] == 1 and summary["average"] == 100
    undone = client.post(f"/rooms/{room_id}/courses/{cid}/progress", headers=SARA, json={"item_id": f"{lid}:notes", "done": False}).json()
    assert undone["percent"] == 50
    assert client.get(f"/rooms/{room_id}/courses/{cid}/progress").status_code == 401


def test_course_access_rules(client):
    room_id = room_with_student(client)
    other_teacher = new_account(client, "teacher", "t2@x.com")
    outsider = new_account(client, "student", "s2@x.com")
    mine, theirs = upload(client), upload(client, headers=other_teacher, data=png_bytes(50, 50))
    cid = client.post("/teacher/courses", headers=MONA, json={"title": "Course"}).json()["course_id"]
    client.post(f"/teacher/courses/{cid}/lessons", headers=MONA, json={"title": "L1", "asset_ids": [mine["asset_id"]]})

    assert client.post(f"/teacher/courses/{cid}/lessons", headers=MONA,
                       json={"title": "L2", "asset_ids": [theirs["asset_id"]]}).status_code == 422       # someone else's file
    assert client.get(f"/teacher/courses/{cid}", headers=other_teacher).status_code == 404
    assert client.get("/teacher/courses", headers=other_teacher).json()["courses"] == []
    foreign_room = client.post("/rooms", headers=other_teacher, json={"name": "Theirs"}).json()["room"]["room_id"]
    assert client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": foreign_room}).status_code == 403
    client.post(f"/teacher/courses/{cid}/publish", headers=MONA, json={"room_id": room_id})
    assert client.get(f"/rooms/{room_id}/courses", headers=outsider).status_code == 403                    # not in the room
    assert client.get(f"/rooms/{room_id}/courses").status_code == 401
    assert client.get("/teacher/courses", headers=SARA).status_code == 403                                  # students use the room routes


def test_lessons_can_be_reordered_and_removed(client):
    cid = client.post("/teacher/courses", headers=MONA, json={"title": "Order"}).json()["course_id"]
    for name in ("A", "B", "C"):
        client.post(f"/teacher/courses/{cid}/lessons", headers=MONA, json={"title": f"Lesson {name}"})
    lessons = client.get(f"/teacher/courses/{cid}", headers=MONA).json()["lessons"]
    first = lessons[0]["lesson_id"]
    moved = client.post(f"/teacher/courses/{cid}/lessons/{first}/move", headers=MONA, json={"direction": "down"}).json()
    assert [l["title"] for l in moved["lessons"]] == ["Lesson B", "Lesson A", "Lesson C"]
    assert client.post(f"/teacher/courses/{cid}/lessons/{first}/move", headers=MONA, json={"direction": "sideways"}).status_code == 422
    left = client.delete(f"/teacher/courses/{cid}/lessons/{first}", headers=MONA).json()
    assert [l["title"] for l in left["lessons"]] == ["Lesson B", "Lesson C"]
    assert client.delete(f"/teacher/courses/{cid}", headers=MONA).status_code == 204
    assert client.get(f"/teacher/courses/{cid}", headers=MONA).status_code == 404


# =============================================================== the web pages, end to end
def login(app, email, password="rafeeq-demo"):
    from starlette.testclient import TestClient

    t = TestClient(app)
    page = t.get("/app/login")
    r = t.post("/app/login", data={"csrf_token": csrf_of(page.text), "email": email, "password": password})
    assert r.status_code == 200 and "/login" not in str(r.url), r.text[:300]
    return t


def csrf_of(html):
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert m, "no CSRF token on the page"
    return m.group(1)


def test_each_studio_is_its_own_page_and_the_sidebar_links_them(client):
    t = login(client.app, "mona@rafeeq.demo")
    pages = {"/app/video": "اكتشاف الفيديو", "/app/presentations": "العروض التقديمية", "/app/media": "المونتاج الذكي",
             "/app/story": "استوديو فيديو القصة", "/app/learning": "منصتي التعليمية", "/app/library": "مكتبتي"}
    for path, heading in pages.items():
        r = t.get(path)
        assert r.status_code == 200 and heading in r.text, path
        assert r.text.count('aria-current="page"') >= 1, path
        for other in pages:                                   # every page carries the whole navigation
            if other != "/app/presentations":                 # Presentations is a tab of the Lesson studio, not a sidebar studio
                assert f'href="{other}' in r.text, (path, other)
        assert f'href="/app/presentations' not in r.text.split('class="main"')[0], path    # not in the sidebar any more
    deck_page = t.get("/app/presentations")                  # the Lesson studio's two tabs, with Presentations current
    assert 'class="subtabs"' in deck_page.text and 'class="subtab active" href="/app/presentations' in deck_page.text
    assert t.get("/app/studio").status_code == 200           # no rooms yet: the rooms page with a hint, never an error
    student = login(client.app, "sara@rafeeq.demo")
    for path in pages:
        assert student.get(path).status_code == 403, path    # students never see the teacher studios
    assert 'href="/app/video' not in student.get("/app/rooms").text


def test_work_flows_between_the_studios_and_persists(client, monkeypatch):
    t = login(client.app, "mona@rafeeq.demo")
    page = t.get("/app/rooms")
    made = t.post("/app/rooms/create", data={"csrf_token": csrf_of(page.text), "name": "Biology 10"})
    room_id = re.search(r"/app/rooms/(r_[0-9a-f]+)/", str(made.url)).group(1)
    code = re.search(r'join-code-value" dir="ltr">([A-Z0-9-]+)<', made.text).group(1)
    deck = make_deck(client, room_id=room_id)
    deck_id = deck["deck_id"]
    studio = t.get(f"/app/rooms/{room_id}/studio")
    assert "ابحث عن شرح بالفيديو" not in studio.text and f"/app/video?room={room_id}" in studio.text   # Lesson Studio stays focused

    # Presentation Studio: list -> editor -> save text
    assert deck_id in t.get(f"/app/presentations?room={room_id}").text
    editor = t.get(f"/app/presentations/{deck_id}")
    assert editor.status_code == 200 and 'name="title_0"' in editor.text and "مقدمة" in editor.text
    token = csrf_of(editor.text)
    saved = t.post(f"/app/presentations/{deck_id}", data={
        "csrf_token": token, "title": "الخلية (منقّح)", "theme_id": "preset:teal", "title_0": "مقدمة جديدة",
        "bullets_0": "نقطة أولى\nنقطة ثانية\n", "notes_0": "ملاحظة", "title_1": "الخاتمة", "bullets_1": "ج", "notes_1": ""})
    assert "حُفظ العرض" in saved.text
    stored = client.get(f"/teacher/decks/{deck_id}", headers=MONA).json()
    assert stored["title"] == "الخلية (منقّح)" and stored["theme_id"] == "preset:teal"
    assert stored["slides"][0]["title"] == "مقدمة جديدة" and stored["slides"][0]["bullets"] == ["نقطة أولى", "نقطة ثانية"]

    # Media Studio: upload -> edit creates a version -> the picture goes on a slide of the deck
    media_token = csrf_of(t.get("/app/media").text)
    up = t.post("/app/media/upload", data={"csrf_token": media_token, "title": "صورة الخلية"},
                files={"file": ("cell.png", png_bytes(), "image/png")})
    asset_id = re.search(r"asset=(a_[0-9a-f]+)", str(up.url)).group(1)
    assert "تم الرفع" in up.text and f"/app/media/{asset_id}/edit" in up.text
    edit = t.post(f"/app/media/{asset_id}/edit", data={"csrf_token": media_token, "rotate": "90", "title": "صورة مدوّرة"})
    version_id = re.search(r"asset=(a_[0-9a-f]+)", str(edit.url)).group(1)
    assert version_id != asset_id and "حُفظت النسخة الجديدة" in edit.text
    nothing = t.post(f"/app/media/{asset_id}/edit", data={"csrf_token": media_token})
    assert "لم تغيّر شيئًا" in nothing.text

    picker = t.get(f"/app/presentations/{deck_id}?asset={version_id}")
    assert "صورة مدوّرة" in picker.text                                                              # the library image is offered
    attached = t.post(f"/app/presentations/{deck_id}/slides/0/image", data={"csrf_token": media_token, "asset_id": version_id})
    assert "تم تحديث صورة الشريحة" in attached.text and f"/app/assets/{version_id}/file" in attached.text
    exported = t.post(f"/app/presentations/{deck_id}/export", data={"csrf_token": media_token})
    assert "اكتمل التصدير" in exported.text and "/file?download=1" in exported.text

    # My Learning Platform: course -> lesson with library files -> publish -> the student sees it
    created = t.post("/app/learning/courses", data={"csrf_token": media_token, "title": "دورة الأحياء", "room_id": room_id})
    cid = re.search(r"/learning/(c_[0-9a-f]+)", str(created.url)).group(1)
    pending_add = t.get(f"/app/learning/{cid}?asset={version_id}")
    assert "صورة مدوّرة" in pending_add.text
    t.post(f"/app/learning/{cid}/lessons", data={"csrf_token": media_token, "title": "الدرس الأول", "notes": "اقرأ أولًا"})
    lesson_id = re.search(r'id="lesson-(l_[0-9a-f]+)"', t.get(f"/app/learning/{cid}").text).group(1)
    t.post(f"/app/learning/{cid}/lessons/{lesson_id}/assets", data={"csrf_token": media_token, "add": version_id})
    published = t.post(f"/app/learning/{cid}/publish", data={"csrf_token": media_token, "room_id": room_id})
    assert "نُشرت الدورة" in published.text

    s = login(client.app, "sara@rafeeq.demo")
    s.post("/app/rooms/join", data={"csrf_token": csrf_of(s.get("/app/rooms").text), "code": code})
    learn = s.get(f"/app/rooms/{room_id}/learn")                                       # the catalog: one card per course
    assert learn.status_code == 200 and "دورة الأحياء" in learn.text and f"/app/rooms/{room_id}/learn/{cid}" in learn.text
    player = s.get(f"/app/rooms/{room_id}/learn/{cid}")                                 # the player opens on the first item
    assert player.status_code == 200 and "الدرس الأول" in player.text and "اقرأ أولًا" in player.text
    image_item = f"{lesson_id}:{version_id}"
    shown = s.get(f"/app/rooms/{room_id}/learn/{cid}?item={image_item}")
    file_url = re.search(r'src="(/app/rooms/[^"]+/file)"', shown.text).group(1)
    assert s.get(file_url).status_code == 200 and s.get(file_url).headers["content-type"].startswith("image/")
    marked = s.post(f"/app/rooms/{room_id}/learn/{cid}/progress",                      # mark complete -> back to the player
                    data={"csrf_token": csrf_of(shown.text), "item_id": image_item, "done": "1"})
    assert marked.status_code == 200 and "lp-item active done" in marked.text                 # the last item: stays, now done
    after = s.get(f"/app/rooms/{room_id}/learn")
    assert "%" in after.text                                                              # the card shows the progress
    preview = t.get(f"/app/learning/{cid}/preview")                                       # the teacher sees the student view
    assert preview.status_code == 200 and "الدرس الأول" in preview.text

    # My Library shows the projects and the new assets, with links to use them elsewhere
    library = t.get("/app/library")
    assert f"/app/presentations/{deck_id}" in library.text and f"/app/learning/{cid}" in library.text
    assert f"/app/media?asset={version_id}" in library.text and "صورة مدوّرة" in library.text


def test_a_finished_lesson_offers_the_other_studios():
    from rafeeq.web.common import templates

    html = templates.get_template("partials/session.html").render(
        t=lambda k, **kw: k, u=lambda p, **q: "/app" + p + ("?" + "&".join(f"{a}={b}" for a, b in q.items()) if q else ""),
        icon=templates.env.globals["icon"], csrf="x", room={"room_id": "r_1", "name": "Physics"},
        session={"session_id": "s_9", "topic": "الضوء", "status": "done", "draft": {"slides": [{"title": "x", "bullets": []}]},
                 "selected_agents": [], "files": [], "revisions_left": 0, "memory": []})
    for path in ("/app/presentations/open?session=s_9", "/app/video?session=s_9", "/app/media?session=s_9", "/app/learning?session=s_9"):
        assert path in html, path
    assert "room=r_1" in html
    waiting = templates.get_template("partials/session.html").render(
        t=lambda k, **kw: k, u=lambda p, **q: "/app" + p, icon=templates.env.globals["icon"], csrf="x",
        room={"room_id": "r_1", "name": "Physics"},
        session={"session_id": "s_9", "topic": "x", "status": "awaiting_approval", "draft": {}, "selected_agents": [],
                 "files": [], "revisions_left": 1, "memory": []})
    assert "/app/presentations/open" not in waiting                                                     # only after approval


def test_youtube_results_page_is_read_without_the_unmaintained_library():
    import json
    from rafeeq.teacher.lecturer import youtube_discovery as yd

    data = {"contents": [{"videoRenderer": {"videoId": "abc12345678", "title": {"runs": [{"text": "Cell"}, {"text": " division"}]},
                                            "ownerText": {"runs": [{"text": "Channel"}]},
                                            "thumbnail": {"thumbnails": [{"url": "small"}, {"url": "big"}]}}},
                         {"videoRenderer": {"videoId": "abc12345678"}},               # duplicate is dropped
                         {"videoRenderer": {"videoId": "zzz12345678"}}]}
    html = "<script>var ytInitialData = " + json.dumps(data) + ";</script>"
    found = yd.parse_results_page(html)
    assert [c["video_id"] for c in found] == ["abc12345678", "zzz12345678"]
    assert found[0]["title"] == "Cell division" and found[0]["thumbnail"] == "big" and found[0]["channel"] == "Channel"
    assert yd.parse_results_page("<html>consent page</html>") == []


def test_blocked_or_captionless_videos_are_not_asked_for_again(monkeypatch):
    import sys, types
    from rafeeq.teacher.lecturer import youtube_discovery as yd

    class IpBlocked(Exception): pass
    class TranscriptsDisabled(Exception): pass
    calls = []

    class Api:
        def fetch(self, video_id, languages):
            calls.append(video_id)
            raise {"a": TranscriptsDisabled, "b": IpBlocked}[video_id]()

    monkeypatch.setitem(sys.modules, "youtube_transcript_api", types.SimpleNamespace(YouTubeTranscriptApi=Api))
    monkeypatch.setattr(yd, "_NO_CAPTIONS", {})
    monkeypatch.setattr(yd, "_BLOCKED_UNTIL", 0.0)
    assert yd.fetch_transcript("a") is None and yd.fetch_transcript("a") is None
    assert calls == ["a"]                                  # captions off: asked once
    assert yd.fetch_transcript("b") is None
    assert yd.fetch_transcript("c") is None                # blocked: nothing else is requested for a while
    assert calls == ["a", "b"]


def test_video_lengths_and_search_fallback_helpers():
    from rafeeq.teacher.lecturer import youtube_discovery as yd

    assert yd.parse_duration("12:34") == 754 and yd.parse_duration("1:02:03") == 3723
    assert yd.parse_duration("") == 0 and yd.parse_duration("LIVE") == 0
    assert yd.expand_queries("ما هي الخلية")[0] == "ما هي الخلية"          # the question itself always comes first


def test_gemini_outline_is_cleaned_and_failures_are_remembered(monkeypatch):
    from rafeeq.teacher.lecturer import youtube_discovery as yd

    monkeypatch.setenv("GOOGLE_API_KEY", "x")
    monkeypatch.setattr(yd, "_NO_OUTLINE", set())
    rows = [{"start": 0, "end": 60, "text": "intro"}, {"start": 60, "end": 500, "text": "long"},
            {"start": 700, "end": 760, "text": "after the end"}, {"start": "bad"}, {"start": 90, "end": 80, "text": "backwards"}]
    monkeypatch.setattr(yd, "_json_from_gemini", lambda *a, **k: rows)
    chunks = yd.outline_with_gemini("abc12345678", 300)
    assert [c["text"] for c in chunks] == ["intro", "long"] and chunks[1]["end_sec"] == 300

    def boom(*a, **k):
        raise RuntimeError("quota")
    monkeypatch.setattr(yd, "_json_from_gemini", boom)
    assert yd.outline_with_gemini("zzz12345678", 300) == []
    monkeypatch.setattr(yd, "_json_from_gemini", lambda *a, **k: rows)
    assert yd.outline_with_gemini("zzz12345678", 300) == []                  # not retried


# ===================================================================== multi-scene story video
def _fake_clip_bytes(color="red", secs=2) -> bytes:
    import shutil
    import subprocess
    import tempfile
    from pathlib import Path

    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg is not installed")
    with tempfile.TemporaryDirectory() as tmp:
        dst = Path(tmp) / "c.mp4"
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"color=c={color}:s=320x180:d={secs}:r=24",
                        "-c:v", "libx264", "-pix_fmt", "yuv420p", str(dst)], check=True)
        return dst.read_bytes()


class _Llm:
    def invoke(self, prompt):
        scenes = ", ".join('{"visual": "a girl %d waters a plant", "narration": "جملة %d"}' % (i, i) for i in range(1, 5))
        return SimpleNamespace(content='<think>hmm</think>```json\n{"character": "girl in a yellow coat", "setting": "green garden", "scenes": [' + scenes + ']}\n```')


def test_story_prompts_repeat_the_character_style_and_setting():
    from rafeeq.teacher import story_video as sv

    plan = sv.plan_story("how plants grow", scenes=3, language="ar", style="flat", llm=_Llm())
    assert len(plan["scenes"]) == 3 and plan["character"] == "girl in a yellow coat"           # extra scenes dropped, <think> ignored
    for i in range(3):
        p = sv.scene_prompt(plan, i, continuity="the girl stands by the tree" if i else "")
        assert "girl in a yellow coat" in p and "green garden" in p and "flat vector" in p.lower() and len(p) <= 800
    assert "previous scene ended" in sv.scene_prompt(plan, 1, "the girl stands by the tree")
    with pytest.raises(sv.StoryError):
        sv.plan_story("x", llm=_Llm())
    with pytest.raises(sv.StoryError):
        sv.plan_story("how plants grow", llm=SimpleNamespace(invoke=lambda p: "no json here"))


def test_story_filter_joins_clips_with_fades_or_cuts():
    from rafeeq.teacher import story_video as sv

    graph, label = sv.stitch_filter([5, 5, 5], 1280, 720, 0.5)
    assert label == "[out]" and graph.count("xfade") == 2 and "offset=4.500" in graph and "offset=9.000" in graph
    cut, _ = sv.stitch_filter([5, 5], 1280, 720, 0)
    assert "concat=n=2" in cut and "xfade" not in cut


def test_story_runs_scene_by_scene_and_ends_as_one_library_video(store, folders, monkeypatch):
    from rafeeq.teacher import assets
    from rafeeq.teacher import story_video as sv
    from rafeeq.teacher.lecturer import video_gen

    monkeypatch.setenv("DASHSCOPE_API_KEY", "x")
    sent, state = [], {"n": 0, "running": set()}

    def submit(prompt, image_bytes=None, image_filename="", duration=5, **kw):
        state["n"] += 1
        sent.append({"prompt": prompt, "image": image_bytes, "duration": duration, **kw})
        return {"ok": True, "task_id": f"t{state['n']}", "model": "wan-fake"}

    def poll(task_id):
        if task_id not in state["running"]:                    # the first poll of each scene: still rendering
            state["running"].add(task_id)
            return {"status": "RUNNING", "video_url": None, "error": None}
        return {"status": "SUCCEEDED", "video_url": "https://x/" + task_id + ".mp4", "error": None}

    clip = _fake_clip_bytes()
    monkeypatch.setattr(video_gen, "submit_task", submit)
    monkeypatch.setattr(video_gen, "poll_task", poll)
    monkeypatch.setattr(video_gen, "download_video", lambda url: clip)

    a = sv.start_story(store, "u_mona", "how plants grow", scenes=3, style="clay", fade=True, llm=_Llm())
    assert a["status"] == "pending" and a["kind"] == "generated_video"
    assert sv.start_story(store, "u_mona", "how plants grow", scenes=3, style="clay", fade=True, llm=_Llm())["asset_id"] == a["asset_id"]   # no second payment

    for _ in range(40):
        a = sv.advance(store, "u_mona", a["asset_id"])
        if a["status"] != "pending":
            break
    assert a["status"] == "ready", a.get("error")
    assert len(sent) == 3 and sent[0]["image"] is None                       # scene 1: text only
    assert sent[1]["image"] and sent[2]["image"]                             # scenes 2 and 3 start from the previous last frame
    assert all("girl in a yellow coat" in s["prompt"] for s in sent)         # same character every time
    assert len({s["seed"] for s in sent}) == 1 and all(s["prompt_extend"] is False for s in sent)   # same seed, prompt left as written
    final = assets.asset_file(a)
    assert final is not None and final.name == "story.mp4" and final.stat().st_size > 1000
    seconds, _, _ = sv.probe(final)
    assert 4.5 < seconds < 6.2                                               # 3 clips of 2 s, 0.5 s fades
    view = sv.story_view(a)
    assert view["scenes_done"] == 3 and view["file_url"] == f"/teacher/assets/{a['asset_id']}/file"
    assert sv.advance(store, "u_mona", a["asset_id"])["status"] == "ready"      # a finished story is left alone
    assert assets.get_asset(store, "u_other", a["asset_id"]) is None


def test_a_failed_scene_stops_the_story_and_retry_restarts_it(store, folders, monkeypatch):
    from rafeeq.teacher import story_video as sv
    from rafeeq.teacher.lecturer import video_gen

    monkeypatch.setenv("DASHSCOPE_API_KEY", "x")
    monkeypatch.setattr(video_gen, "submit_task", lambda *a, **k: {"ok": False, "error": "quota exceeded"})
    a = sv.start_story(store, "u_mona", "how plants grow", scenes=2, llm=_Llm())
    a = sv.advance(store, "u_mona", a["asset_id"])
    assert a["status"] == "failed" and "quota exceeded" in a["error"] and "Scene 1" in a["error"]
    assert sv.start_story(store, "u_mona", "how plants grow", scenes=2, llm=_Llm())["status"] == "failed"      # no silent restart
    again = sv.start_story(store, "u_mona", "how plants grow", scenes=2, llm=_Llm(), retry=True)
    assert again["asset_id"] == a["asset_id"] and again["status"] == "pending"


def test_a_video_can_become_a_gif(tmp_path):
    from rafeeq.teacher import story_video as sv

    src = tmp_path / "v.mp4"
    src.write_bytes(_fake_clip_bytes(secs=3))
    gif = sv.to_gif(src, tmp_path / "v.gif", start=0, end=2)
    assert gif.read_bytes()[:6] == b"GIF89a"


def test_story_video_api_and_page(client, monkeypatch):
    from rafeeq.teacher import media_tools
    from rafeeq.teacher import story_video as sv
    from rafeeq.teacher.lecturer import video_gen

    body = {"topic": "how plants grow", "scenes": 2, "style": "flat", "language": "en", "fade": True}
    assert client.post("/teacher/stories", json=body).status_code == 401
    assert client.post("/teacher/stories", headers=SARA, json=body).status_code == 403
    monkeypatch.setattr(media_tools, "ffmpeg_path", lambda: None)
    assert client.post("/teacher/stories", headers=MONA, json=body).status_code == 503          # no ffmpeg: say so before paying
    monkeypatch.setattr(media_tools, "ffmpeg_path", lambda: "/usr/bin/ffmpeg")
    monkeypatch.delenv("DASHSCOPE_API_KEY", raising=False)
    assert client.post("/teacher/stories", headers=MONA, json=body).status_code == 503          # no key
    monkeypatch.setenv("DASHSCOPE_API_KEY", "x")
    assert client.post("/teacher/stories", headers=MONA, json={**body, "scenes": 99}).status_code == 422

    monkeypatch.setattr(sv, "plan_story", lambda topic, scenes, language, style, llm=None: {
        "style": style, "character": "girl", "setting": "garden",
        "scenes": [{"visual": "a", "narration": "n1"}, {"visual": "b", "narration": "n2"}]})
    submitted = []
    monkeypatch.setattr(video_gen, "submit_task", lambda *a, **k: submitted.append(1) or {"ok": True, "task_id": f"t{len(submitted)}", "model": "m"})
    monkeypatch.setattr(video_gen, "poll_task", lambda task_id: {"status": "RUNNING", "video_url": None, "error": None})
    first = client.post("/teacher/stories", headers=MONA, json=body)
    assert first.status_code == 201 and first.json()["status"] == "pending" and first.json()["scene_count"] == 2
    asset_id = first.json()["asset_id"]
    assert client.post("/teacher/stories", headers=MONA, json=body).json()["asset_id"] == asset_id   # the same request is the same story
    status = client.get(f"/teacher/stories/{asset_id}", headers=MONA)
    assert status.status_code == 200 and status.json()["scenes"][0]["state"] == "running" and len(submitted) == 1
    assert client.get(f"/teacher/stories/{asset_id}", headers=MONA).json()["scenes"][0]["state"] == "running" and len(submitted) == 1
    assert client.get(f"/teacher/stories/{asset_id}", headers=SARA).status_code == 403
    assert client.get("/teacher/stories/a_nope", headers=MONA).status_code == 404
    other = new_account(client, "teacher", "other@x.com")
    assert client.get(f"/teacher/stories/{asset_id}", headers=other).status_code == 404            # not their story

    t = login(client.app, "mona@rafeeq.demo")
    page = t.get("/app/story")
    assert page.status_code == 200 and "how plants grow" in page.text and 'name="scenes"' in page.text
    poll = t.get(f"/app/story/{asset_id}")
    assert poll.status_code == 200 and "hx-get" in poll.text and "n1" in poll.text


def test_a_video_becomes_a_gif_that_can_go_on_a_slide(client):
    import shutil

    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg is not installed")
    video = client.post("/teacher/media/upload", headers=MONA, files={"file": ("v.mp4", _fake_clip_bytes(secs=2), "video/mp4")})
    assert video.status_code == 201, video.text
    gif = client.post(f"/teacher/media/{video.json()['asset_id']}/edit", headers=MONA, json={"ops": {"to_gif": True}, "title": "Loop"})
    assert gif.status_code == 201, gif.text
    gif = gif.json()
    assert gif["kind"] == "image" and gif["origin"] == "edited" and gif["parent_asset_id"] == video.json()["asset_id"]
    assert client.get(gif["file_url"], headers=MONA).content[:6] == b"GIF89a"
    again = client.post(f"/teacher/media/{gif['asset_id']}/edit", headers=MONA, json={"ops": {"rotate": 90}})
    assert again.status_code == 422                                                            # editing would flatten the animation
    deck = make_deck(client)
    placed = client.post(f"/teacher/decks/{deck['deck_id']}/slides/0/image", headers=MONA, json={"asset_id": gif["asset_id"]})
    assert placed.status_code == 200 and placed.json()["slides"][0]["image_asset_id"] == gif["asset_id"]
    exported = client.post(f"/teacher/decks/{deck['deck_id']}/export", headers=MONA)
    assert exported.status_code in (200, 201), exported.text
    data = client.get(exported.json()["file_url"], headers=MONA).content
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert any(n.startswith("ppt/media/") and n.endswith(".gif") for n in z.namelist())     # the animated GIF travels inside the .pptx


def test_story_planner_asks_for_a_teaching_arc_not_a_character_walk():
    from rafeeq.teacher import story_video as sv

    seen = {}

    class Llm(_Llm):
        def invoke(self, prompt):
            seen["prompt"] = prompt
            return super().invoke(prompt)

    sv.plan_story("دورة الماء", 4, "ar", "flat", llm=Llm())
    p = seen["prompt"]
    assert "TEACH" in p and "recap" in p and "REAL thing" in p
    assert "empty string" in p                                    # a process may have no character


def test_story_narration_is_one_voice_request_cut_into_scenes(tmp_path, monkeypatch):
    import shutil
    import subprocess

    from rafeeq.teacher import media_tools
    from rafeeq.teacher import story_video as sv

    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg is not installed")
    clips = []
    for i in range(2):
        c = tmp_path / f"c{i}.mp4"
        c.write_bytes(_fake_clip_bytes("red" if i == 0 else "blue", 2))
        clips.append(c)
    asked = []

    def fake_speech(text, voice, dst):
        asked.append(text)
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "sine=f=300:d=6", str(dst)], check=True)
        return dst

    monkeypatch.setattr(media_tools, "synthesize_speech", fake_speech)
    story = {"voice": "ar-EG-SalmaNeural", "fade": False, "scenes": [{"narration": "a" * 50}, {"narration": "b" * 10}]}
    out, mix = sv._narrate(story, clips, tmp_path)
    assert len(asked) == 1 and "a" * 50 in asked[0] and "b" * 10 in asked[0]       # one request for the whole video
    assert sv.probe(out[0])[0] >= 5.3                                               # 5 of the 6 voice seconds belong to scene 1
    assert abs(sv.probe(out[1])[0] - 2.0) < 0.2                                     # 1 s fits in its 2 s scene
    assert sv.probe(mix)[0] >= sv.probe(out[0])[0]


def test_voice_is_cut_at_pauses_near_the_proportional_point():
    from rafeeq.teacher import story_video as sv

    assert sv.split_points([50, 10], 6.0, []) == [5.0, 6.0]                         # no pause: proportional
    assert sv.split_points([50, 10], 6.0, [4.6, 1.0]) == [4.6, 6.0]                 # snaps to the pause 0.4 s away
    assert sv.split_points([50, 10], 6.0, [1.0]) == [5.0, 6.0]                      # a pause too far is ignored
    assert sv.split_points([10, 10, 10], 9.0, []) == [3.0, 6.0, 9.0]


def test_story_planner_accepts_a_reply_made_of_blocks():
    from rafeeq.teacher import story_video as sv

    class Llm:
        def invoke(self, prompt):
            body = '{"character": "", "setting": "bright", "scenes": [{"visual": "water heats up", "narration": "أ"}, {"visual": "steam rises", "narration": "ب"}]}'
            return SimpleNamespace(content=[{"type": "text", "text": body[:40]}, {"type": "text", "text": body[40:]}])

    plan = sv.plan_story("دورة الماء", 2, "ar", "flat", llm=Llm())
    assert len(plan["scenes"]) == 2 and plan["character"] == ""


def test_story_planner_writes_the_subject_notes_first_and_uses_them():
    from rafeeq.teacher import story_video as sv

    prompts = []

    class Llm:
        def invoke(self, prompt):
            prompts.append(prompt)
            if "Write the teaching notes" in prompt:
                return SimpleNamespace(content='{"objective": "o", "facts": ["الماء يتبخر عند 100 درجة", "البخار يبرد فيصير سحابًا", "السحاب يمطر"], "analogy": "كوب شاي"}')
            return SimpleNamespace(content='{"character": "", "setting": "s", "scenes": [{"visual": "kettle boils", "narration": "أ"}, {"visual": "cloud forms", "narration": "ب"}]}')

    sv.plan_story("دورة الماء", 2, "ar", "flat", llm=Llm())
    assert len(prompts) == 2 and "Write the teaching notes" in prompts[0]
    assert "الماء يتبخر عند 100 درجة" in prompts[1] and "كوب شاي" in prompts[1]


def test_story_planner_still_works_when_the_notes_step_fails():
    from rafeeq.teacher import story_video as sv

    class Llm(_Llm):
        def invoke(self, prompt):
            if "Write the teaching notes" in prompt:
                raise RuntimeError("boom")
            return super().invoke(prompt)

    assert len(sv.plan_story("الخلية", 3, "ar", "flat", llm=Llm())["scenes"]) == 3


def test_speech_uses_gemini_voice_and_falls_back_to_edge(tmp_path, monkeypatch):
    import base64
    import shutil

    import httpx

    from rafeeq.teacher import media_tools

    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg is not installed")
    monkeypatch.setenv("GOOGLE_API_KEY", "k")
    monkeypatch.delenv("TTS_PROVIDER", raising=False)
    pcm = b"\x00\x10" * 24000                                           # one second of 24 kHz 16-bit mono
    calls = []

    def post(url, headers=None, json=None, timeout=None):
        calls.append((url, json))
        body = {"candidates": [{"content": {"parts": [{"inlineData": {"mimeType": "audio/L16;rate=24000", "data": base64.b64encode(pcm).decode()}}]}}]}
        return httpx.Response(200, json=body, request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", post)
    out = media_tools.synthesize_speech("مرحبا بكم", "ar-EG-SalmaNeural", tmp_path / "a.mp3")
    assert out.stat().st_size > 0 and "gemini" in calls[0][0]
    assert calls[0][1]["generationConfig"]["speechConfig"]["voiceConfig"]["prebuiltVoiceConfig"]["voiceName"] == "Kore"

    # Gemini refuses -> the Microsoft voice is tried, and its answer is used
    monkeypatch.setattr(httpx, "post", lambda url, **kw: httpx.Response(403, request=httpx.Request("POST", url)))
    def edge_ok(text, voice, dst):
        dst.write_bytes(b"ID3")
        return dst

    monkeypatch.setattr(media_tools, "_edge_speech", edge_ok)
    assert media_tools.synthesize_speech("مرحبا بكم", "ar-EG-ShakirNeural", tmp_path / "b.mp3").read_bytes() == b"ID3"

    # both fail -> one message with both reasons
    def edge_fail(text, voice, dst):
        raise media_tools.MediaError("Microsoft said no")

    monkeypatch.setattr(media_tools, "_edge_speech", edge_fail)
    with pytest.raises(media_tools.MediaError) as e:
        media_tools.synthesize_speech("مرحبا بكم", "ar-EG-SalmaNeural", tmp_path / "c.mp3")
    assert "403" in str(e.value) and "Microsoft said no" in str(e.value)


def test_story_finish_runs_in_the_background_and_does_not_start_twice(monkeypatch):
    import threading

    from rafeeq.teacher import assets
    from rafeeq.teacher import story_video as sv

    release, started = threading.Event(), []
    state = {"status": "pending"}

    def slow_finish(store, teacher_id, asset, story):
        started.append(1)
        release.wait(5)
        state["status"] = "ready"

    monkeypatch.setattr(sv, "_finish", slow_finish)
    monkeypatch.setattr(sv, "FINISH_WAIT_SECONDS", 0.2)
    monkeypatch.setattr(assets, "get_asset", lambda store, tid, aid: {"asset_id": aid, "status": state["status"], "settings": {}})
    asset = {"asset_id": "a_bg", "status": "pending", "settings": {}}

    first = sv._finish_in_background(None, "t", asset, {})
    assert first["status"] == "pending" and "a_bg" in sv._finishing          # answered at once while it keeps working
    sv._finish_in_background(None, "t", asset, {})                           # the next poll does not start a second one
    assert started == [1]
    release.set()
    for _ in range(50):
        if "a_bg" not in sv._finishing:
            break
        threading.Event().wait(0.05)
    assert state["status"] == "ready" and "a_bg" not in sv._finishing


def test_gemini_voice_waits_and_retries_when_google_says_too_many_requests(tmp_path, monkeypatch):
    import base64
    import shutil

    import httpx

    from rafeeq.teacher import media_tools

    if shutil.which("ffmpeg") is None:
        pytest.skip("ffmpeg is not installed")
    monkeypatch.setenv("GOOGLE_API_KEY", "k")
    monkeypatch.setenv("TTS_PROVIDER", "gemini")
    waited, answers = [], []
    pcm = base64.b64encode(b"\x00\x10" * 24000).decode()

    def post(url, headers=None, json=None, timeout=None):
        answers.append(1)
        req = httpx.Request("POST", url)
        if len(answers) < 3:
            return httpx.Response(429, text='{"error": {"details": [{"retryDelay": "7s"}]}}', request=req)
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"inlineData": {"data": pcm}}]}}]}, request=req)

    monkeypatch.setattr(httpx, "post", post)
    monkeypatch.setattr(media_tools.time, "sleep", waited.append)
    assert media_tools.synthesize_speech("مرحبا بكم", "ar-EG-SalmaNeural", tmp_path / "a.mp3").stat().st_size > 0
    assert waited == [8.0, 8.0]                                       # Google's 7 s + 1 s, twice

    # it never retries forever, and it says what to do
    answers.clear()
    monkeypatch.setattr(httpx, "post", lambda url, **kw: httpx.Response(429, text="{}", request=httpx.Request("POST", url)))
    with pytest.raises(media_tools.MediaError) as e:
        media_tools.synthesize_speech("مرحبا بكم", "ar-EG-SalmaNeural", tmp_path / "b.mp3")
    assert "429" in str(e.value) and "quota" in str(e.value)
