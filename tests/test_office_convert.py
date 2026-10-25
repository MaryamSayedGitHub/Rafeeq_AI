"""PowerPoint -> PDF with LibreOffice. The real conversion runs only where soffice is installed."""
import pytest

from rafeeq.knowledge.ingestion import office


def test_extensions():
    assert office.is_convertible("a/Deck.PPTX") and office.is_convertible("x.ppt") and not office.is_convertible("x.pdf")


def test_missing_libreoffice_gives_a_clear_message(monkeypatch, tmp_path):
    f = tmp_path / "d.pptx"
    f.write_bytes(b"x")
    monkeypatch.setattr(office, "find_soffice", lambda: None)
    with pytest.raises(office.ConversionError) as e:
        office.to_pdf(f)
    assert "LibreOffice" in str(e.value)


def test_soffice_path_env_is_used(monkeypatch, tmp_path):
    exe = tmp_path / "soffice"
    exe.write_text("")
    monkeypatch.setenv("SOFFICE_PATH", str(exe))
    assert office.find_soffice() == str(exe)


def test_real_conversion(tmp_path):
    if office.find_soffice() is None:
        pytest.skip("LibreOffice is not installed")
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Gradient descent"
    s.placeholders[1].text = "Follow the slope of the loss"
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Second slide"
    src = tmp_path / "lesson.pptx"
    prs.save(src)
    pdf = office.to_pdf(src, tmp_path / "out")
    assert pdf.name == "lesson.pdf" and pdf.read_bytes().startswith(b"%PDF")
